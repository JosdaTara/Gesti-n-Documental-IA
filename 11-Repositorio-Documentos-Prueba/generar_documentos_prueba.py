# -*- coding: utf-8 -*-
"""Genera el repositorio de documentos de prueba de SIGAD (carpeta 11).

Crea 36 documentos distribuidos en 5 categorías y 4 formatos (.txt, .md,
.pdf, .docx) con contenido empresarial realista en español, listos para
cargarse en el sistema y validar la clasificación IA y la búsqueda semántica.

Requisitos: fpdf2 y python-docx instalados.

Uso:
    python generar_documentos_prueba.py
"""

import random
from pathlib import Path

from fpdf import FPDF
from docx import Document

BASE = Path(__file__).resolve().parent

PROVEEDORES_F = [
    ("Alimentos del Valle S.A.", "900.123.456-7"),
    ("Papeles y Empaques Ltda.", "830.987.654-3"),
    ("Ferreterías del Eje S.A.", "900.456.789-1"),
    ("Distribuciones Andinas SAS", "901.234.567-8"),
    ("Textiles Nacionales Cía.", "860.345.678-9"),
    ("Químicos del Pacífico SA", "890.112.233-4"),
    ("Logística Express de Bogotá", "901.876.543-2"),
    ("Maquinaria Industrial Ltda.", "800.987.123-5"),
]

DESTINATARIOS = [
    "Bodega Central Bogotá",
    "Sucursal Cali",
    "Centro de Distribución Medellín",
    "Punto de Venta Barranquilla",
    "Almacén Bucaramanga",
    "Sucursal Pereira",
    "Centro Logístico Cartagena",
]

TRANSPORTADORAS = [
    "Envíos Andinos",
    "Transportes Córdoba S.A.",
    "Logística Veloz",
    "Rapibogotá Express",
    "Interandina de Carga",
]

MONEDA = "COP"


def numero(n: int, d: int = 0) -> str:
    return f"{n:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def text_factura(i: int) -> str:
    prov, nit = PROVEEDORES_F[i % len(PROVEEDORES_F)]
    base = random.randint(150000, 9000000)
    iva = round(base * 0.19)
    total = base + iva
    return (
        f"FACTURA DE VENTA FE-{1000 + i:04d}\n"
        f"NIT proveedor: {nit}\n"
        f"Proveedor: {prov}\n"
        f"Fecha de emisión: {10 + i:02d}/0{i % 6 + 1}/2026\n"
        f"Ciudad: Bogotá D.C.\n"
        f"Subtotal: {numero(base)} {MONEDA}\n"
        f"IVA (19%): {numero(iva)} {MONEDA}\n"
        f"Total a pagar: {numero(total)} {MONEDA}\n"
        f"Número de factura: FACT-2026-{142 + i * 7:04d}\n"
        f"Medio de pago: transferencia bancaria\n"
        f"Condición: 30 días\n"
        f"Forma de pago: contado/crédito\n\n"
        f"Resumen de ítems:\n"
        f"- Producto A x {2 + i}: {numero(base // 2)} {MONEDA}\n"
        f"- Producto B x {1 + i % 3}: {numero(base - base // 2)} {MONEDA}\n"
        f"Bienes entregados al cliente conforme a la orden de compra asociada."
    )


def text_guia(i: int) -> str:
    dest = DESTINATARIOS[i % len(DESTINATARIOS)]
    trans = TRANSPORTADORAS[i % len(TRANSPORTADORAS)]
    return (
        f"GUÍA DE DESPACHO GD-{3000 + i}\n"
        f"Remitente: Bodega Central Bogotá\n"
        f"Destinatario: {dest}\n"
        f"Transportadora: {trans}\n"
        f"Número de caso: DI-{800 + i}\n"
        f"Fecha de despacho: {5 + i:02d}/0{i % 6 + 1}/2026\n"
        f"Bultos: {8 + i}  Peso total: {100 + i * 25} kg\n"
        f"Referencia: factura FACT-2026-{142 + i * 7:04d}\n"
        f"Observaciones: mercancía en buen estado, empaque sellado."
    )


def text_orden(i: int) -> str:
    prov, nit = PROVEEDORES_F[(i + 3) % len(PROVEEDORES_F)]
    cant = 50 + i * 8
    precio = 30000 + i * 3500
    total = cant * precio
    return (
        f"ORDEN DE COMPRA OC-{500 + i}\n"
        f"Proveedor: {prov}\n"
        f"NIT: {nit}\n"
        f"Fecha de solicitud: {15 + i:02d}/0{i % 6 + 1}/2026\n"
        f"Cantidad: {cant}  Precio unitario: {numero(precio)} {MONEDA}\n"
        f"Total de la orden: {numero(total)} {MONEDA}\n"
        f"Condiciones de entrega: {10 + i * 2} días\n"
        f"Términos de pago: contado a la entrega\n"
        f"Cotización de referencia: COT-{8800 + i}\n"
        f"Libre de fletes. Aprobada por el área de compras."
    )


def text_contrato(i: int) -> str:
    prov, nit = PROVEEDORES_F[(i + 5) % len(PROVEEDORES_F)]
    return (
        f"CONTRATO DE PRESTACIÓN DE SERVICIOS CT-{200 + i}\n"
        f"Entre Distribuidora Andina S.A.S. y {prov} (NIT {nit})\n"
        f"Objeto: suministro e implementación de servicios documentales.\n"
        f"Cláusula 1: vigencia de {12 + i * 6} meses.\n"
        f"Cláusula 2: obligaciones de las partes.\n"
        f"Cláusula 3: términos de pago y facturación mensual.\n"
        f"Cláusula 4: confidencialidad de la información.\n"
        f"Cláusula 5: resolución de conflictos por arbitraje.\n"
        f"Firmas y legalización en notaría {8 + i} del círculo de Bogotá.\n"
        f"Valor total del contrato: {numero(2000000 + i * 450000)} {MONEDA}"
    )


def text_acta(i: int) -> str:
    dest = DESTINATARIOS[(i + 2) % len(DESTINATARIOS)]
    return (
        f"ACTA DE RECEPCIÓN AR-{700 + i}\n"
        f"Entrega de mercancía en {dest}.\n"
        f"Fecha: {20 + i:02d}/0{i % 6 + 1}/2026\n"
        f"Se recibe a satisfacción la totalidad de los bienes relacionados\n"
        f"con la guía de despacho GD-{3000 + i}.\n"
        f"Observaciones: recepción conforme, sin novedades.\n"
        f"Conformidad del área de bodega y del cliente final."
    )


def escribir_docx(path: Path, titulo: str, texto: str) -> None:
    doc = Document()
    doc.add_heading(titulo, level=1)
    for linea in texto.splitlines():
        if linea.strip():
            doc.add_paragraph(linea)
    doc.save(str(path))


def escribir_pdf(path: Path, titulo: str, texto: str) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", "B", 14)
    pdf.multi_cell(0, 8, titulo, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("helvetica", size=10)
    for linea in texto.splitlines():
        if linea.strip():
            pdf.multi_cell(0, 6, linea, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.ln(2)
    pdf.output(str(path))


def main() -> None:
    random.seed(2026)
    plantillas = [
        ("FACTURA", 8, text_factura, "FACTURA-FE"),
        ("GUIA_DESPACHO", 7, text_guia, "GUIA-DESPACHO-GD"),
        ("ORDEN_COMPRA", 7, text_orden, "ORDEN-COMPRA-OC"),
        ("CONTRATO", 7, text_contrato, "CONTRATO-CT"),
        ("ACTA_RECEPCION", 7, text_acta, "ACTA-RECEPCION-AR"),
    ]
    formatos = ["txt", "md", "pdf", "docx"]
    total = 0
    for categoria, cantidad, fn, prefijo in plantillas:
        carpeta = BASE / categoria
        carpeta.mkdir(exist_ok=True)
        for i in range(cantidad):
            texto = fn(i)
            titulo = texto.splitlines()[0]
            ext = formatos[i % len(formatos)]
            archivo = carpeta / f"{prefijo}-{i + 1:03d}.{ext}"
            if ext == "pdf":
                escribir_pdf(archivo, titulo, texto)
            elif ext == "docx":
                escribir_docx(archivo, titulo, texto)
            else:
                archivo.write_text(texto, encoding="utf-8")
            total += 1
    print(f"Generados {total} documentos de prueba en {BASE}")


if __name__ == "__main__":
    main()